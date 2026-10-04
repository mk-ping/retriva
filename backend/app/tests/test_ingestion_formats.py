import ingestion


def test_new_types_are_accepted():
    for name, ext in [("a.PPTX", ".pptx"), ("b.xlsx", ".xlsx"), ("c.html", ".html"), ("d.htm", ".htm")]:
        assert ingestion.validate_file(name, 100) == ext


def test_html_scripts_are_stripped(tmp_path):
    p = tmp_path / "page.html"
    p.write_text("<html><body><script>alert(1)</script><h1>Title</h1><p>Hello world</p></body></html>", encoding="utf-8")
    text = " ".join(c["text"] for c in ingestion.extract_and_chunk(str(p), ".html"))
    assert "Hello world" in text and "alert" not in text


def test_xlsx_rows_keep_their_column_names(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["name", "city"])
    ws.append(["Anna", "Delhi"])
    ws.append(["Ravi", "Pune"])
    p = tmp_path / "d.xlsx"
    wb.save(str(p))
    text = " ".join(c["text"] for c in ingestion.extract_and_chunk(str(p), ".xlsx"))
    assert "name: Anna" in text and "city: Pune" in text and "People" in text


def test_pptx_page_is_the_slide_number(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    for title in ("First slide", "Second slide"):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = title
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    pages = {c["page"]: c["text"] for c in ingestion.extract_and_chunk(str(p), ".pptx")}
    assert "First slide" in pages[1] and "Second slide" in pages[2]