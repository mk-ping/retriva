import os
import re

from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader, TextLoader, CSVLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import ALLOWED_EXTENSIONS, MAX_FILE_SIZE_MB, OCR_MAX_PAGES

# fixed-size chunks keep every piece small enough for the embedding model to read fully
# and keep the prompt well under Groq's free-tier tokens-per-minute limit
_splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)


class IngestionError(Exception):
    pass


def validate_file(filename: str, size_bytes: int):
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise IngestionError(f"Unsupported file type: {ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")
    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if size_bytes > max_bytes:
        raise IngestionError(f"File exceeds the {MAX_FILE_SIZE_MB}MB limit.")
    return ext


def _split(text: str, page):
    return [{"text": p, "page": page} for p in _splitter.split_text(text) if p.strip()]


def _ocr_pdf(filepath: str) -> list[dict]:
    try:
        import pytesseract
        from pdf2image import convert_from_path, pdfinfo_from_path
    except ImportError:
        raise IngestionError("OCR is not installed on the server.")

    total = pdfinfo_from_path(filepath)["Pages"]
    chunks = []
    for page_no in range(1, min(total, OCR_MAX_PAGES) + 1):
        images = convert_from_path(filepath, dpi=200, first_page=page_no, last_page=page_no)
        text = pytesseract.image_to_string(images[0])
        chunks += _split(text, page_no)
    return chunks


def _pdf_chunks(filepath: str) -> list[dict]:
    docs = PyPDFLoader(filepath).load()
    chunks = []
    for d in docs:
        if not d.page_content.strip():
            continue
        page = d.metadata.get("page")  # PyPDFLoader pages are 0-based
        chunks += _split(d.page_content, page + 1 if isinstance(page, int) else None)

    # almost no embedded text means a scanned PDF, so fall back to OCR
    total_chars = sum(len(d.page_content.strip()) for d in docs)
    if total_chars >= 20 * max(1, len(docs)):
        return chunks
    return _ocr_pdf(filepath) or chunks


def _pptx_chunks(filepath: str) -> list[dict]:
    from pptx import Presentation

    chunks = []
    for number, slide in enumerate(Presentation(filepath).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            nf = slide.notes_slide.notes_text_frame
            if nf is not None and nf.text.strip():
                parts.append("Speaker notes: " + nf.text.strip())
        if parts:
            chunks += _split("\n".join(parts), number)  # page = slide number
    return chunks


def _xlsx_chunks(filepath: str) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(filepath, read_only=True, data_only=True)
    chunks = []
    try:
        for ws in wb.worksheets:
            prefix = f"[Sheet: {ws.title}]\n"
            header, buf, size = None, [], 0

            def flush():
                nonlocal buf, size
                if buf:
                    chunks.append({"text": prefix + "\n".join(buf), "page": None})
                buf, size = [], 0

            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n > 50000:
                    break
                cells = ["" if c is None else str(c).strip() for c in row]
                if not any(cells):
                    continue
                if header is None:
                    header = cells
                    continue
                pairs = []
                for i, v in enumerate(cells):
                    if v:
                        name = header[i] if i < len(header) and header[i] else f"col{i + 1}"
                        pairs.append(f"{name}: {v}")
                line = "; ".join(pairs)
                if size + len(line) > 800:
                    flush()
                buf.append(line)
                size += len(line)
            flush()
    finally:
        wb.close()
    return chunks


def _html_chunks(filepath: str) -> list[dict]:
    from bs4 import BeautifulSoup

    with open(filepath, "rb") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "svg"]):
        tag.decompose()
    text = re.sub(r"\n\s*\n+", "\n\n", soup.get_text("\n")).strip()
    return _split(text, None)


def extract_and_chunk(filepath: str, ext: str) -> list[dict]:
    """Returns a list of {"text": str, "page": int | None}."""
    if ext == ".pdf":
        return _pdf_chunks(filepath)
    if ext == ".pptx":
        return _pptx_chunks(filepath)
    if ext == ".xlsx":
        return _xlsx_chunks(filepath)
    if ext in (".html", ".htm"):
        return _html_chunks(filepath)
    if ext == ".csv":
        # CSV rows stay whole, because splitting a row by length would separate its fields
        return [{"text": d.page_content, "page": None} for d in CSVLoader(filepath).load() if d.page_content.strip()]
    if ext == ".docx":
        docs = Docx2txtLoader(filepath).load()
    elif ext in (".txt", ".md"):
        docs = TextLoader(filepath, encoding="utf-8").load()
    else:
        raise IngestionError(f"No loader for {ext}")

    chunks = []
    for d in docs:
        if d.page_content.strip():
            chunks += _split(d.page_content, None)
    return chunks